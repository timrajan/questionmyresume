"""Extract plain text from an uploaded resume (PDF, TXT, DOC, DOCX)."""

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .config import ALLOWED_EXTENSIONS

# Magic bytes, so a renamed file can't masquerade as an allowed type.
_SIGNATURES = {
    ".pdf": [b"%PDF"],
    ".docx": [b"PK\x03\x04"],
    ".doc": [b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"],
}

MIN_RESUME_CHARS = 100


class ResumeParseError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported file type '{ext or 'none'}'. Upload a PDF, TXT, DOC or DOCX file."
        )
    if not data:
        raise ResumeParseError("The uploaded file is empty.")

    signatures = _SIGNATURES.get(ext)
    if signatures and not any(data.startswith(sig) for sig in signatures):
        raise ResumeParseError(f"This file doesn't look like a real {ext} file.")

    try:
        if ext == ".pdf":
            text = _from_pdf(data)
        elif ext == ".docx":
            text = _from_docx(data)
        elif ext == ".doc":
            text = _from_doc(data)
        else:
            text = _from_txt(data)
    except ResumeParseError:
        raise
    except Exception as exc:  # parser libraries raise a zoo of exception types
        raise ResumeParseError(f"Couldn't read the {ext} file: {exc}") from exc

    text = _normalize(text)
    if len(text) < MIN_RESUME_CHARS:
        raise ResumeParseError(
            "Couldn't find enough text in the resume. If it's a scanned PDF, "
            "export it as a text-based PDF or DOCX and try again."
        )
    return text


def _from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Many resume templates put content in tables.
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_doc(data: bytes) -> str:
    """Legacy Word .doc: use antiword if installed, else macOS textutil."""
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(data)
        path = tmp.name
    try:
        if shutil.which("antiword"):
            cmd = ["antiword", path]
        elif shutil.which("textutil"):
            cmd = ["textutil", "-convert", "txt", "-stdout", path]
        else:
            raise ResumeParseError(
                "Legacy .doc files need 'antiword' installed on the server. "
                "Save the resume as .docx or PDF instead."
            )
        result = subprocess.run(cmd, capture_output=True, timeout=30, check=True)
        return result.stdout.decode("utf-8", errors="replace")
    finally:
        Path(path).unlink(missing_ok=True)


def _from_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()
