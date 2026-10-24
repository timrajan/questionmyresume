import io

import pytest
from docx import Document

from app.resume_parser import ResumeParseError, extract_text

RESUME = "Jane Doe\nSenior Engineer\n" + "Built data pipelines processing 2TB/day with Python and Spark. " * 3


def test_txt():
    assert "Jane Doe" in extract_text("resume.txt", RESUME.encode())


def test_docx_includes_tables():
    doc = Document()
    doc.add_paragraph(RESUME)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Kubernetes"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("resume.docx", buf.getvalue())
    assert "Jane Doe" in text and "Kubernetes" in text


@pytest.mark.parametrize("name", ["resume.png", "resume.exe", "resume"])
def test_rejects_other_extensions(name):
    with pytest.raises(ResumeParseError, match="Unsupported"):
        extract_text(name, RESUME.encode())


def test_rejects_renamed_file():
    with pytest.raises(ResumeParseError, match="real .pdf"):
        extract_text("resume.pdf", RESUME.encode())


def test_rejects_too_little_text():
    with pytest.raises(ResumeParseError, match="enough text"):
        extract_text("resume.txt", b"hi")
